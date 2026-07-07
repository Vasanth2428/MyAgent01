import docx

doc = docx.Document("UPDATED REPORT.docx")
total = len(doc.paragraphs)
for idx in range(total - 10, total):
    p = doc.paragraphs[idx]
    print(f"P[{idx}] ({p.style.name}): {repr(p.text)}")
