import docx

doc = docx.Document("UPDATED REPORT.docx")

for idx in range(240, 275):
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        print(f"P[{idx}] ({p.style.name}): {repr(p.text)}")
