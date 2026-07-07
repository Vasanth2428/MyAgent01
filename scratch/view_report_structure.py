import docx

doc_path = "UPDATED REPORT.docx"
doc = docx.Document(doc_path)

with open("scratch/report_paragraphs.txt", "w", encoding="utf-8") as f:
    f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
    for idx, para in enumerate(doc.paragraphs):
        style = para.style.name if para.style else "Normal"
        text = para.text.strip()
        f.write(f"Para [{idx}] | Style: {style} | Text: {text[:100]}...\n")

print("Done. Structure written to scratch/report_paragraphs.txt")
