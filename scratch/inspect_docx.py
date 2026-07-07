import os
import docx

doc_path = "UPDATED REPORT.docx"
output_path = "scratch/docx_content.txt"

if not os.path.exists(doc_path):
    print(f"Error: {doc_path} not found.")
    exit(1)

doc = docx.Document(doc_path)
print(f"Document loaded. Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

with open(output_path, "w", encoding="utf-8") as f:
    f.write(f"Document loaded. Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}\n")
    
    f.write("\n=== HEADINGS AND PARAGRAPHS ===\n")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            # Check if it has a heading style
            if para.style.name.startswith("Heading"):
                f.write(f"\n[{para.style.name}] {text}\n")
            else:
                # Only write short samples or titles to keep it readable, but since we are writing to file, write all text
                f.write(f"{text}\n")

print(f"Docx content successfully written to {output_path}")
