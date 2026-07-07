import os
from pptx import Presentation

ppt_path = "UPDATED PPT.pptx"
output_path = "scratch/ppt_content.txt"

if not os.path.exists(ppt_path):
    print(f"Error: {ppt_path} not found.")
    exit(1)

prs = Presentation(ppt_path)
print(f"Presentation loaded. Total slides: {len(prs.slides)}")

with open(output_path, "w", encoding="utf-8") as f:
    f.write(f"Presentation loaded. Total slides: {len(prs.slides)}\n")
    for i, slide in enumerate(prs.slides):
        f.write(f"\n--- Slide {i+1} ---\n")
        # Check for title
        if slide.shapes.title:
            f.write(f"Title: {slide.shapes.title.text}\n")
        else:
            f.write("Title: (No direct title placeholder)\n")
            
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        f.write(f"  - {text}\n")

print(f"Slide content successfully written to {output_path}")
