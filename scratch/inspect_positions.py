import pptx

prs = pptx.Presentation("UPDATED PPT.pptx")

for slide_num in [7, 10]:
    slide = prs.slides[slide_num - 1]
    print(f"\nSlide {slide_num}:")
    for shape in slide.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  Picture: {shape.name}")
            print(f"    Left: {shape.left}, Top: {shape.top}")
            print(f"    Width: {shape.width}, Height: {shape.height}")
        elif shape.has_text_frame:
            print(f"  TextBox: {shape.name} ('{shape.text_frame.text.strip()[:30]}')")
            print(f"    Left: {shape.left}, Top: {shape.top}")
            print(f"    Width: {shape.width}, Height: {shape.height}")
        elif shape.has_table:
            print(f"  Table: {shape.name}")
            print(f"    Left: {shape.left}, Top: {shape.top}")
            print(f"    Width: {shape.width}, Height: {shape.height}")
