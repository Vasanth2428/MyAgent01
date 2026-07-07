from pptx import Presentation

prs = Presentation("UPDATED PPT.pptx")
slide = prs.slides[9] # Slide 10 is index 9

with open("scratch/table_content.txt", "w", encoding="utf-8") as f:
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            f.write(f"Table found: {len(table.rows)} rows x {len(table.columns)} columns\n")
            for r_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                f.write(f"  Row {r_idx}: {cells}\n")
print("Done. Output written to scratch/table_content.txt")
