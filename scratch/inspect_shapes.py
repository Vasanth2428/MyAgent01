import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation("UPDATED PPT.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"\nSlide {idx+1}:")
    for shape in slide.shapes:
        print(f"  Shape Name: {shape.name}, Type: {shape.shape_type}")
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"    -> PICTURE found")
        if shape.has_text_frame:
            print(f"    -> Text: {shape.text_frame.text.strip()[:60]}")
