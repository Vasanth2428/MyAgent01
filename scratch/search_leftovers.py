import docx

doc_path = "UPDATED REPORT.docx"
doc = docx.Document(doc_path)

keywords = ["hela", "tinospora", "cervical", "cancer", "mtt", "docking", "phytochemical", "berberine", "quercetin", "palmatine", "magnoflorine", "tinosporaside"]

found = False

# Search in paragraphs
for idx, para in enumerate(doc.paragraphs):
    text_lower = para.text.lower()
    for kw in keywords:
        if kw in text_lower:
            print(f"Match in Para [{idx}] for keyword '{kw}': {para.text[:120]}...")
            found = True

# Search in tables
for idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text_lower = cell.text.lower()
            for kw in keywords:
                if kw in text_lower:
                    print(f"Match in Table {idx+1}, Row {r_idx}, Cell {c_idx} for keyword '{kw}': {cell.text[:120]}...")
                    found = True

if not found:
    print("Zero leftovers found! The report is 100% clean.")
else:
    print("Leftovers found. We need to clean them up.")
