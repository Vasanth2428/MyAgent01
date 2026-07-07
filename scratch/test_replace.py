import docx

doc_path = "UPDATED REPORT.docx"
doc = docx.Document(doc_path)

print(f"Para 95 before: '{doc.paragraphs[95].text}'")

doc.paragraphs[95].text = "Design and Optimization of an Agentic RAG Context Engine & Multi-Agent Orchestration Platform"

print(f"Para 95 after: '{doc.paragraphs[95].text}'")

doc.save(doc_path)
print("Saved. Loading again to verify...")

doc2 = docx.Document(doc_path)
print(f"Para 95 verified: '{doc2.paragraphs[95].text}'")
