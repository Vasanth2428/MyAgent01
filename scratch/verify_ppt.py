import pptx

prs = pptx.Presentation("UPDATED PPT.pptx")
print(f"Total slides in refined PPT: {len(prs.slides)}")

issues = []

# Biology terms search
bio_kw = ["hela", "tinospora", "cervical cancer", "mtt assay", "phytochemical", 
          "berberine", "quercetin", "palmatine", "magnoflorine", "tinosporaside",
          "formazan", "monolayer", "trypsin"]

for idx, slide in enumerate(prs.slides):
    # Check text shapes
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.lower()
            for kw in bio_kw:
                if kw in t:
                    issues.append(f"Slide {idx+1}: Found biology keyword '{kw}' in shape {shape.name}: {shape.text_frame.text[:80]}")
        
        if shape.has_table:
            # Check table contents
            for r_idx, row in enumerate(shape.table.rows):
                for cell in row.cells:
                    t = cell.text.lower()
                    for kw in bio_kw:
                        if kw in t:
                            issues.append(f"Slide {idx+1}: Found biology keyword '{kw}' in Table 1 R{r_idx}: {cell.text}")

        # Check for remaining pictures on slides other than slide 7 & 10 (our technical figures)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            if idx+1 not in [7, 10]:
                issues.append(f"Slide {idx+1}: Picture shape '{shape.name}' still remains!")

with open("scratch/ppt_verification_report.txt", "w", encoding="utf-8") as f:
    f.write("=== PPT VERIFICATION REPORT ===\n")
    if not issues:
        f.write("STATUS: ALL CLEAN - No issues found in Slide Deck!\n")
    else:
        f.write(f"STATUS: {len(issues)} issues found:\n")
        for issue in issues:
            f.write(f"  - {issue}\n")
            
print("Verification report written to scratch/ppt_verification_report.txt")
