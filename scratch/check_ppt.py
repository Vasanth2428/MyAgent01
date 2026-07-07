import pptx

prs = pptx.Presentation("UPDATED PPT.pptx")
print(f"Total slides: {len(prs.slides)}")

with open("scratch/ppt_audit.txt", "w", encoding="utf-8") as f:
    for s_idx, slide in enumerate(prs.slides):
        f.write(f"\n--- Slide {s_idx+1} ---\n")
        # Check shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                f.write(f"[{shape.name}]: {shape.text_frame.text.strip()}\n")
            if shape.has_table:
                f.write(f"[{shape.name}] (TABLE):\n")
                for r_idx, row in enumerate(shape.table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    f.write(f"  R{r_idx}: {cells}\n")

print("PPT Audit written to scratch/ppt_audit.txt")
