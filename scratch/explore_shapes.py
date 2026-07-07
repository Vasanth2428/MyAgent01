import os
from pptx import Presentation

ppt_path = "UPDATED PPT.pptx"
output_path = "scratch/shapes_outline.txt"

if not os.path.exists(ppt_path):
    print(f"Error: {ppt_path} not found.")
    exit(1)

prs = Presentation(ppt_path)

with open(output_path, "w", encoding="utf-8") as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n=========================================\n")
        f.write(f"SLIDE {idx+1}\n")
        f.write(f"=========================================\n")
        
        # Check for title
        if slide.shapes.title:
            f.write(f"Title Shape Name: {slide.shapes.title.name}, Text: {slide.shapes.title.text}\n")
        else:
            f.write("No direct title placeholder\n")
            
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"\n  Shape [{s_idx}]: {shape.name} | Type: {shape.shape_type}\n")
            if shape.has_text_frame:
                f.write(f"    Text Frame Paragraphs count: {len(shape.text_frame.paragraphs)}\n")
                for p_idx, para in enumerate(shape.text_frame.paragraphs):
                    f.write(f"      Para [{p_idx}]: {para.text.strip()}\n")

print(f"Done exploring shapes. Output written to {output_path}")
