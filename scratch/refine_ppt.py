"""
Refine the slide deck (UPDATED PPT.pptx) to match the technical content of the report.
Replaces all biology-themed placeholders with RAG & Multi-Agent platform details,
removes biological pictures, and inserts the generated technical diagrams.
"""

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

ppt_path = "UPDATED PPT.pptx"
prs = pptx.Presentation(ppt_path)

print(f"Loaded PPT presentation. Total slides: {len(prs.slides)}")

def clean_and_format_text(shape, text, font_name="Calibri", font_size=Pt(16), bold=False, italic=False, color=None, alignment=None):
    """Safely replace text in a shape, applying clean formatting."""
    if not shape.has_text_frame:
        return
    
    tf = shape.text_frame
    tf.clear()
    
    # Split by newline and add paragraphs
    lines = text.split('\n')
    for idx, line in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.text = line
        p.font.name = font_name
        p.font.size = font_size
        p.font.bold = bold
        p.font.italic = italic
        
        if color:
            p.font.color.rgb = color
            
        if alignment:
            p.alignment = alignment
            
        # Add space after paragraphs
        p.space_after = Pt(6)

def delete_all_pictures_in_slide(slide):
    """Deletes all picture shapes in a given slide to clear biological elements."""
    pics = []
    for shape in slide.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pics.append(shape)
            
    for pic in pics:
        # We delete the shape using the underlying XML element tree
        sp = pic._element
        sp.getparent().remove(sp)
    print(f"  Deleted {len(pics)} pictures.")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1: Title Slide
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 1...")
slide1 = prs.slides[0]
delete_all_pictures_in_slide(slide1)

for shape in slide1.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "COE INTERNSHIP (JUNE - 2026)", font_size=Pt(36), bold=True)
    elif shape.name == "Subtitle 2":
        clean_and_format_text(shape, "Department of Computer Science and Engineering\nChennai Institute of Technology", font_size=Pt(18), italic=True)
    elif shape.name == "TextBox 5":
        info_text = (
            "NAME: Vasanth S\n"
            "REG NO: 210424104085\n"
            "SEMESTER: V\n"
            "FACULTY MENTOR: Dr. S. Pavithra M.E, Ph.D.\n"
            "TITLE: Design and Optimization of an Agentic RAG Context Engine & Multi-Agent Orchestration Platform"
        )
        clean_and_format_text(shape, info_text, font_size=Pt(14))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2: Introduction
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 2...")
slide2 = prs.slides[1]
delete_all_pictures_in_slide(slide2)

for shape in slide2.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "INTRODUCTION", font_size=Pt(28), bold=True)
    elif shape.name == "Rectangle 3":
        intro_text = (
            "In autonomous software development, monolithic LLM developer prompts suffer from context window limits, high latency, and expensive API costs.\n\n"
            "This internship focuses on designing a modular, supervisor-directed RAG framework that isolates agent concerns into a collaborative state graph using LangGraph.\n\n"
            "We introduce surgical context reduction tools, dynamic hybrid retrieval, and semantic query caching to optimize token load and inference latency."
        )
        clean_and_format_text(shape, intro_text, font_size=Pt(15))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3: Goals of the Internship
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 3...")
slide3 = prs.slides[2]
delete_all_pictures_in_slide(slide3)

for shape in slide3.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Goals of the Internship", font_size=Pt(28), bold=True)
    elif shape.name == "Content Placeholder 2":
        goals_text = (
            "• Implement a production-grade multi-agent coordination system using LangGraph and a Supervisor routing pattern.\n"
            "• Configure a robust hybrid search index using Weaviate Cloud (BM25 + Semantic) with local fallback databases.\n"
            "• Design surgical context compression, symbol-based signature fetchers, and diff-based code generators (apply_surgical_edit).\n"
            "• Debug and resolve live API routing issues, caching crashes, and safety boundary constraints on a multi-repo codebase."
        )
        clean_and_format_text(shape, goals_text, font_size=Pt(15))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4: Why Modular RAG & Multi-Agents?
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 4...")
slide4 = prs.slides[3]
delete_all_pictures_in_slide(slide4)

for shape in slide4.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Why Modular RAG & Multi-Agents?", font_size=Pt(28), bold=True)
    elif shape.name == "Rectangle 2":
        why_text = (
            "Feeding entire code files and log streams to LLMs results in 'lost in the middle' problems, poor code generation, and high costs.\n\n"
            "Isolating complex tasks (retrieval, search, code writing, execution, criticism) into specialized worker agents improves system reliability.\n\n"
            "By adopting a surgical diff-based editing tool (apply_surgical_edit), we target only the modified functions, bypassing the need to regenerate unaltered boilerplate code, which reduces latency from ~30s to <1s."
        )
        clean_and_format_text(shape, why_text, font_size=Pt(15))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5: Methodology: Multi-Agent Architecture
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 5...")
slide5 = prs.slides[4]
delete_all_pictures_in_slide(slide5)

for shape in slide5.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Methodology: Multi-Agent Architecture", font_size=Pt(28), bold=True)
    elif shape.name == "Content Placeholder 3":
        meth_text = (
            "• Supervisor Pattern: Coordinates routing using a central LangGraph state. Dispatches subtasks and resumes state upon worker outputs.\n"
            "• RAG Worker: Queries indexed knowledge via Weaviate and returns grounded document matches.\n"
            "• Coding & Critic Workers: The Coding Worker generates structural modifications; the Critic Worker verifies code safety and syntactical accuracy.\n"
            "• Web & Scraper Workers: Fetch live web data (via Tavily API) and extract raw text from target links.\n"
            "• Utility Worker: Computes statistics and manages session history."
        )
        clean_and_format_text(shape, meth_text, font_size=Pt(14))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 6: Methodology: Core Retrieval Innovations
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 6...")
slide6 = prs.slides[5]
delete_all_pictures_in_slide(slide6)

for shape in slide6.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Methodology: Core Retrieval Innovations", font_size=Pt(28), bold=True)
    elif shape.name == "TextBox 8":
        ret_text = (
            "• Dynamic Hybrid Retrieval (Alpha Shifts): Automatically detects query type. For technical/code queries, shifts alpha to keyword-matching (0.25); for natural language, shifts to semantic (0.75).\n"
            "• Semantic Cache Layer: An in-memory, query-hash keyed cache matching previous queries to bypass vector DB calls entirely, eliminating latency for repeated questions.\n"
            "• Neural Reranking: Employs a local FlashRank Reranker (ms-marco-MiniLM-L-6-v2) to filter initial search outputs and surface high-relevance chunks.\n"
            "• Hybrid Local Database Fallback: Uses a local JSON/SQLite vector store when cloud connectivity is degraded, ensuring zero uptime dependency."
        )
        clean_and_format_text(shape, ret_text, font_size=Pt(14))
    elif shape.name == "Rectangle 2":
        # Make this rectangle shape empty and transparent or small
        clean_and_format_text(shape, "")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 7: Visuals: RAG & Orchestration Engine Flow
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 7...")
slide7 = prs.slides[6]
delete_all_pictures_in_slide(slide7)

for shape in slide7.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Visuals: RAG & Orchestration Engine Flow", font_size=Pt(28), bold=True)
    elif shape.name == "TextBox 4":
        clean_and_format_text(shape, "Figure 2. Retrieval & Generation Pipeline Flow", font_size=Pt(11), italic=True, alignment=PP_ALIGN.CENTER)
        shape.left = Inches(0.5)
        shape.top = Inches(5.6)
        shape.width = Inches(5.0)
        shape.height = Inches(0.5)
    elif shape.name == "TextBox 11":
        clean_and_format_text(shape, "Figure 3. System Latency Heatmap (ms)", font_size=Pt(11), italic=True, alignment=PP_ALIGN.CENTER)
        shape.left = Inches(6.2)
        shape.top = Inches(5.6)
        shape.width = Inches(5.0)
        shape.height = Inches(0.5)
    elif shape.name == "TextBox 26":
        clean_and_format_text(shape, "")

# Insert the new figures
if os.path.exists("scratch/fig3_rag_pipeline.png"):
    slide7.shapes.add_picture("scratch/fig3_rag_pipeline.png", Inches(0.5), Inches(1.8), Inches(5.0), Inches(3.6))
    print("  Inserted fig3_rag_pipeline.png")
if os.path.exists("scratch/fig5_heatmap.png"):
    slide7.shapes.add_picture("scratch/fig5_heatmap.png", Inches(6.2), Inches(1.8), Inches(5.0), Inches(3.6))
    print("  Inserted fig5_heatmap.png")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 8: Performance Savings: Surgical Code Edits
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 8...")
slide8 = prs.slides[7]
delete_all_pictures_in_slide(slide8)

for shape in slide8.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Performance Savings: Surgical Code Edits", font_size=Pt(28), bold=True)
    elif shape.name == "TextBox 40":
        surg_text = (
            "• The Problem: Conventional coding agents rewrite entire code files (7,000+ tokens) to change a few lines of code, consuming high generation tokens and increasing timeout rates.\n\n"
            "• Surgical Approach: Implemented apply_surgical_edit which requires the agent to output only search-and-replace target blocks.\n\n"
            "• AST-Based Header Fetching: Uses a parser (fetch_file_headers) to extract classes, functions, and import structures without loading the full implementation.\n\n"
            "• Token Truncation: Truncates verbose command line outputs and logs to fit strict context parameters (achieving up to 88.3% token savings)."
        )
        clean_and_format_text(shape, surg_text, font_size=Pt(14))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 9: Performance Benchmarks & Metrics
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 9...")
slide9 = prs.slides[8]
delete_all_pictures_in_slide(slide9)

for shape in slide9.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Performance Benchmarks & Metrics", font_size=Pt(28), bold=True)
    elif shape.name == "TextBox 3":
        bench_text = (
            "• Structural Header Scan: Reduced context overhead by 99.79% (a 7,138-token file reduced to a 15-token header outline).\n\n"
            "• Surgical Generation: Reduced output token volume by 99.55% during codebase editing.\n\n"
            "• Semantic Context Compression: Compressed raw retrieval contexts by 20% to 58.3% while dynamically retaining critical query terms.\n\n"
            "• Log Output Summarization: Truncated long grep logs and compile outputs, achieving 88.3% token savings and preventing context bloating."
        )
        clean_and_format_text(shape, bench_text, font_size=Pt(14))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 10: Visuals: Performance Tables
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 10...")
slide10 = prs.slides[9]
delete_all_pictures_in_slide(slide10)

# Replace table content
for shape in slide10.shapes:
    if shape.name == "TextBox 2":
        clean_and_format_text(shape, "Visuals: Performance Tables", font_size=Pt(28), bold=True)
    elif shape.name == "Table 1":
        table = shape.table
        print(f"  Found table with {len(table.rows)} rows and {len(table.columns)} columns.")
        
        # We have 5 columns: 'Optimization Scenario', 'Conventional Method', 'Optimized Method', 'Token Savings (%)', 'Key Advantage'
        headers = ['Optimization Scenario', 'Conventional Method', 'Optimized Method', 'Token Savings (%)', 'Key Advantage']
        
        rows_data = [
            ['Imports & Signatures Scan', 'Read full file (7,138 tokens)', 'fetch_file_headers (15 tokens)', '99.79%', 'Excludes implementation noise'],
            ['Viewing Specific Function', 'Read full file (7,138 tokens)', 'get_pruned_context (808 tokens)', '88.68%', 'Isolates target code blocks'],
            ['Applying Code Edits', 'Rewrite full file (7,138 tokens)', 'apply_surgical_edit (32 tokens)', '99.55%', 'Avoids rewriting full files'],
            ['Log Output Parsing', 'Raw console logs (1,800 tokens)', 'Line-level truncation (211 tokens)', '88.30%', 'Retains critical compile errors'],
            ['Full Site Scaffolding', 'Traditional LLM flow (2,210 tokens)', 'Optimized agentic loop (115 tokens)', '94.80%', 'Orchestrator-driven scaffolding']
        ]
        
        # Let's resize columns to make it look clean
        table.columns[0].width = Inches(1.8)
        table.columns[1].width = Inches(1.3)
        table.columns[2].width = Inches(1.3)
        table.columns[3].width = Inches(1.0)
        table.columns[4].width = Inches(1.8)
        
        # Populate headers
        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = h
            cell.text_frame.paragraphs[0].font.name = "Calibri"
            cell.text_frame.paragraphs[0].font.size = Pt(10)
            cell.text_frame.paragraphs[0].font.bold = True
            
        # Populate data rows
        for row_idx, r_data in enumerate(rows_data):
            for col_idx, val in enumerate(r_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = val
                cell.text_frame.paragraphs[0].font.name = "Calibri"
                cell.text_frame.paragraphs[0].font.size = Pt(9)
                
        # Resize table position to fit nicely
        shape.left = Inches(0.5)
        shape.top = Inches(1.3)
        shape.width = Inches(7.2)
        shape.height = Inches(3.0)

    elif shape.name == "TextBox 9":
        caption_text = "Figure 4. Token load comparison between conventional read/write and optimized surgical coding pipelines."
        clean_and_format_text(shape, caption_text, font_size=Pt(11), italic=True, alignment=PP_ALIGN.CENTER)
        shape.left = Inches(8.0)
        shape.top = Inches(5.6)
        shape.width = Inches(4.5)
        shape.height = Inches(0.8)

# Insert the token savings chart next to the table
if os.path.exists("scratch/fig2_token_savings.png"):
    slide10.shapes.add_picture("scratch/fig2_token_savings.png", Inches(8.0), Inches(1.5), Inches(4.5), Inches(3.8))
    print("  Inserted fig2_token_savings.png")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 11: Conclusion
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 11...")
slide11 = prs.slides[10]
delete_all_pictures_in_slide(slide11)

for shape in slide11.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "Conclusion", font_size=Pt(28), bold=True)
    elif shape.name == "Content Placeholder 2":
        conclusion_text = (
            "• Combining LangGraph multi-agent coordination with surgical context compression makes agentic development pipelines highly viable, fast, and cost-effective.\n\n"
            "• Architectural Insights: Semantic caches and AST-based signature lookups are essential to prevent context window bloating in production-grade software engineering agents.\n\n"
            "• The hybrid Weaviate vector engine with local fallback ensures high retrieval accuracy and system availability, showing that token-optimized context engineering is critical for scaling autonomous agentic pipelines."
        )
        clean_and_format_text(shape, conclusion_text, font_size=Pt(15))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 12: OUTCOME
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 12...")
slide12 = prs.slides[11]
delete_all_pictures_in_slide(slide12)

for shape in slide12.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "OUTCOME", font_size=Pt(28), bold=True)
    elif shape.name == "Content Placeholder 6":
        outcome_text = (
            "• Design and orchestration of multi-agent state machines (LangGraph) with routing supervisor pattern.\n\n"
            "• Core integration of vector database endpoints (Weaviate Cloud) and dynamic alpha shifting (0.25 / 0.75).\n\n"
            "• Advanced context pruning, token metrics, caching, and safety guardrails.\n\n"
            "• Developed professional academic-level benchmarking and verification procedures."
        )
        clean_and_format_text(shape, outcome_text, font_size=Pt(15))
    elif shape.name == "TextBox 14":
        clean_and_format_text(shape, "")

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 13: Thank You
# ─────────────────────────────────────────────────────────────────────────────
print("\nRefining Slide 13...")
slide13 = prs.slides[12]
delete_all_pictures_in_slide(slide13)

for shape in slide13.shapes:
    if shape.name == "Title 1":
        clean_and_format_text(shape, "THANK YOU", font_size=Pt(36), bold=True, alignment=PP_ALIGN.CENTER)
        # Center the shape on the slide
        shape.left = Inches(1.5)
        shape.top = Inches(2.0)
        shape.width = Inches(10.3)
        shape.height = Inches(1.5)

# Add presenter info box on slide 13
presenter_box = slide13.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.3), Inches(2.0))
presenter_text = (
    "Presented by: Vasanth S (Reg no: 210424104085)\n"
    "Department of Computer Science and Engineering\n"
    "Chennai Institute of Technology"
)
clean_and_format_text(presenter_box, presenter_text, font_size=Pt(16), alignment=PP_ALIGN.CENTER)

# Save the presentation
prs.save(ppt_path)
print("\n✓ PPT refinement completed successfully!")
