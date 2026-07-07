import docx

doc_path = "UPDATED REPORT.docx"
doc = docx.Document(doc_path)

with open("scratch/tables_outline.txt", "w", encoding="utf-8") as f:
    f.write(f"Total tables: {len(doc.tables)}\n")
    for idx, table in enumerate(doc.tables):
        f.write(f"\nTable {idx+1}: {len(table.rows)} rows x {len(table.columns)} columns\n")
        for r_idx in range(len(table.rows)):
            cells = [cell.text.strip().replace("\n", " ") for cell in table.rows[r_idx].cells]
            f.write(f"  Row {r_idx}: {cells}\n")

print("Done. Output written to scratch/tables_outline.txt")
